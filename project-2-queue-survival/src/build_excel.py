"""
Step 5 - Assemble every result into a single reviewable Excel workbook.

The workbook is the deliverable for anyone who does not want to run Python. It
leads with method and caveats rather than burying them, because the two data
traps in this dataset (competing risks, and regions that do not report operation
dates) are the difference between the right answer and a confident wrong one.
"""
import pathlib
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Repo root, derived from this file's own location so the project runs
# anywhere after a clone rather than only on the machine that wrote it.
ROOT = str(pathlib.Path(__file__).resolve().parents[2]).replace("\\", "/")
OUT = ROOT + "/project-2-queue-survival/output"
PROC = ROOT + "/data/processed"
XLSX = OUT + "/interconnection_queue_analysis.xlsx"

INK = "FF0B0B0B"
HEAD_FILL = PatternFill("solid", fgColor="FF1C5CAB")
BAND_FILL = PatternFill("solid", fgColor="FFF2F5FA")
THIN = Side(style="thin", color="FFD8D7D0")

NOTES = [
    ("Interconnection Queue Survival Analysis", 16, True),
    ("", 11, False),
    ("QUESTION", 12, True),
    ("Of the projects that ask to connect to the US grid, how many are ever built, "
     "how long does it take, and what predicts success?", 11, False),
    ("", 11, False),
    ("SOURCE", 12, True),
    ("Lawrence Berkeley National Laboratory, 'Queued Up' 2026 Edition, data through "
     "year-end 2025. Licensed CC BY 4.0. Compiled by LBNL and GridTracker, covering "
     "7 ISOs/RTOs and 50 utilities (~98% of US generating capacity).", 11, False),
    ("", 11, False),
    ("TWO TRAPS IN THIS DATA, AND HOW THEY ARE HANDLED", 12, True),
    ("1. Withdrawal is a competing risk, not censoring.", 11, True),
    ("   A project leaves the queue either by being built or by being withdrawn. "
     "Standard Kaplan-Meier treats withdrawal as 'we stopped watching', which assumes "
     "those projects might still be built later. They cannot. Using Kaplan-Meier here "
     "overstates the ten-year completion probability by 2.9x (58.0% vs the correct "
     "20.0%). This analysis uses Aalen-Johansen cumulative incidence instead.", 11, False),
    ("2. Four regions do not report commercial operation dates.", 11, True),
    ("   ISO-NE reports zero operation dates across 227 operational projects. West, "
     "NYISO and Southeast report 20-49%. Any completion curve built on those regions "
     "shows near-zero completion, which is a fact about their data, not their grid. "
     "They are excluded from completion estimates. See the 'Data quality' sheet.", 11, False),
    ("", 11, False),
    ("KNOWN LIMITATIONS", 12, True),
    ("- 30% of operational and 35% of withdrawn projects lack an outcome date and are "
     "excluded from duration estimates. Dropped rows are similar to kept rows on size "
     "and cohort year (see missingness audit), so the bias appears mild, but it is not zero.", 11, False),
    ("- Battery and Solar+Battery entered the queue mostly after 2019. Their five-year "
     "estimates rest on a few hundred early projects and will move as cohorts mature.", 11, False),
    ("- The logistic model is associational, not causal. It ranks and forecasts; it does "
     "not prove that filing in ERCOT causes a project to succeed.", 11, False),
    ("", 11, False),
    ("HOW TO REPRODUCE", 12, True),
    ("Run src/profile_data.py, build_dataset.py, survival.py, completion_model.py, "
     "make_charts.py, then build_excel.py. Raw source file is in data/raw/.", 11, False),
]

SHEETS = [
    ("Headline", OUT + "/cif_summary.csv",
     "Probability of each outcome, by years since interconnection request"),
    ("By technology", OUT + "/cif_by_tech.csv",
     "Completion probability by resource type. n_observed_5yr shows how much data backs each estimate."),
    ("By region", OUT + "/cif_by_region.csv",
     "Completion probability by market. Reliable reporters only."),
    ("By era", OUT + "/cif_by_era.csv",
     "Completion within 5 years, by the era the project filed. Like-for-like horizon."),
    ("Model A odds ratios", OUT + "/model_a_odds_ratios.csv",
     "Logistic model using only request-time features. Odds ratio > 1 raises the odds of being built."),
    ("Model B odds ratios", OUT + "/model_b_odds_ratios.csv",
     "Adds the interconnection-agreement milestone. Not usable as a forecast feature."),
    ("Model calibration", OUT + "/model_a_calibration.csv",
     "Predicted vs actual completion rate by decile, held-out set."),
    ("Scenario", OUT + "/scenario_battery_by_region.csv",
     "Identical 200 MW battery project scored in each market."),
    ("Data quality", PROC + "/region_reporting_quality.csv",
     "Share of operational projects with a reported commercial operation date, by region."),
]

CHARTS = [
    ("01_outcome_composition.png", "What happens to a project"),
    ("02_naive_vs_competing_risks.png", "Why the method matters"),
    ("03_completion_by_tech.png", "By technology"),
    ("04_completion_by_region.png", "By market"),
    ("05_era_trend.png", "Getting worse over time"),
    ("06_scenario_by_region.png", "Worked example"),
]


def write_notes(writer):
    pd.DataFrame({"": [n[0] for n in NOTES]}).to_excel(
        writer, sheet_name="Read me first", index=False, header=False)


with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
    write_notes(writer)
    for name, path, _ in SHEETS:
        pd.read_csv(path).to_excel(writer, sheet_name=name, index=False)
    # placeholder sheet so charts have somewhere to live
    pd.DataFrame().to_excel(writer, sheet_name="Charts", index=False)

wb = load_workbook(XLSX)

# --- Read me first ----------------------------------------------------------
ws = wb["Read me first"]
ws.column_dimensions["A"].width = 118
for i, (text, size, bold) in enumerate(NOTES, start=1):
    c = ws.cell(row=i, column=1)
    c.font = Font(size=size, bold=bold, color=INK)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if size == 16:
        ws.row_dimensions[i].height = 26
    elif len(text) > 110:
        ws.row_dimensions[i].height = 46
ws.sheet_view.showGridLines = False

# --- Data sheets ------------------------------------------------------------
for name, _, caption in SHEETS:
    ws = wb[name]
    ws.insert_rows(1, 2)
    ws["A1"] = caption
    ws["A1"].font = Font(size=11, italic=True, color="FF52514E")

    header_row = 3
    ncols = ws.max_column
    for j in range(1, ncols + 1):
        c = ws.cell(row=header_row, column=j)
        c.fill = HEAD_FILL
        c.font = Font(bold=True, color="FFFFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row,
                                         max_col=ncols)):
        for c in row:
            c.border = Border(bottom=THIN)
            if isinstance(c.value, float):
                c.number_format = "0.00"
            if i % 2 == 1:
                c.fill = BAND_FILL

    for j in range(1, ncols + 1):
        longest = max((len(str(ws.cell(row=r, column=j).value or ""))
                       for r in range(header_row, min(ws.max_row, 200) + 1)), default=10)
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 11), 34)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.sheet_view.showGridLines = False

# --- Charts -----------------------------------------------------------------
ws = wb["Charts"]
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
row = 2
for fname, caption in CHARTS:
    path = os.path.join(OUT, fname)
    if not os.path.exists(path):
        continue
    c = ws.cell(row=row, column=2, value=caption)
    c.font = Font(size=13, bold=True, color=INK)
    img = XLImage(path)
    scale = 780 / img.width
    img.width, img.height = 780, int(img.height * scale)
    ws.add_image(img, "B{}".format(row + 1))
    row += int(img.height / 19) + 4

wb.save(XLSX)

print("workbook written ->", XLSX)
print("sheets:", ", ".join(wb.sheetnames))
